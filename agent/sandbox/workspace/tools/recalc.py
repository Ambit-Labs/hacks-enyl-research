#!/usr/bin/env python3
"""Recalculate every formula in an xlsx workbook with headless LibreOffice.

Runs the same `soffice --headless --convert-to` invocation the
SpreadsheetBench evaluator (eval/sb.py) uses, so a value read here matches
what the judge will see. The recalculated file is copied to the requested
output path; pass --cells to also print specific answer values on stdout.

Usage:
    python3 recalc.py <in.xlsx> <out.xlsx> [--cells Sheet1!B6:B8]
"""

import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter

CONVERT_FILTER = "xlsx:Calc MS Excel 2007 XML"


def soffice_path():
    if os.environ.get("SOFFICE"):
        return os.environ["SOFFICE"]
    return shutil.which("soffice") or shutil.which("libreoffice")


def recalculate(xlsx_path, out_path):
    exe = soffice_path()
    if not exe:
        raise RuntimeError("soffice not found on PATH; is libreoffice-calc installed?")
    with tempfile.TemporaryDirectory() as out_dir, tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [
                exe,
                f"-env:UserInstallation={Path(profile).resolve().as_uri()}",
                "--headless",
                "--calc",
                "--convert-to",
                CONVERT_FILTER,
                "--outdir",
                out_dir,
                str(xlsx_path),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=180,
        )
        converted = Path(out_dir) / Path(xlsx_path).name
        shutil.copyfile(converted, out_path)


def parse_ref(ref):
    """Parse a 'Sheet1!B6:B8' style reference into (sheet, min_col, min_row, max_col, max_row)."""
    match = re.match(r"^(?:'?([^'!]+)'?!)?([A-Za-z0-9:$]+)$", ref)
    if not match:
        raise ValueError(f"unrecognized cell reference: {ref}")
    sheet, cell_range = match.groups()
    cell_range = cell_range.replace("$", "")
    if ":" not in cell_range:
        cell_range = f"{cell_range}:{cell_range}"
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    return sheet, min_col, min_row, max_col, max_row


def print_cells(xlsx_path, refs):
    wb = load_workbook(xlsx_path, data_only=True)
    for ref in refs:
        sheet_name, min_col, min_row, max_col, max_row = parse_ref(ref)
        ws = wb[sheet_name] if sheet_name else wb.active
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                coord = f"{get_column_letter(col)}{row}"
                value = ws.cell(row=row, column=col).value
                print(f"{ws.title}!{coord}\t{value}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="path to the source xlsx")
    parser.add_argument("output", help="path to write the recalculated xlsx")
    parser.add_argument(
        "--cells",
        action="append",
        default=[],
        help="cell or range to print after recalculation, e.g. Sheet1!B6:B8 (repeatable)",
    )
    args = parser.parse_args()

    recalculate(args.input, args.output)
    print(f"recalculated {args.input} -> {args.output}")
    if args.cells:
        print_cells(args.output, args.cells)


if __name__ == "__main__":
    try:
        main()
    except subprocess.CalledProcessError as exc:
        print(f"recalc: soffice failed: {exc.stderr}", file=sys.stderr)
        sys.exit(1)
    except Exception as exc:  # noqa: BLE001 - surface any failure to the caller
        print(f"recalc: {exc}", file=sys.stderr)
        sys.exit(1)
