#!/usr/bin/env python3
"""Print every sheet of an xlsx workbook as a compact grid.

Each cell shows its formula (when present) and its last-computed value, so
the model can see both what a cell does and what it currently holds without
opening the file in a spreadsheet program. Wide or tall sheets are
truncated to keep the output readable.

Usage:
    python3 dump_workbook.py <xlsx> [--max-rows N] [--max-cols N]
"""

import argparse
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DEFAULT_MAX_ROWS = 120
DEFAULT_MAX_COLS = 30


def format_cell(formula_cell, value_cell):
    formula = formula_cell.value
    value = value_cell.value
    if isinstance(formula, str) and formula.startswith("="):
        shown_value = "" if value is None else value
        return f"{formula} => {shown_value}"
    if formula is None:
        return ""
    return str(formula)


def dump_sheet(formula_ws, value_ws, max_rows, max_cols):
    rows = min(formula_ws.max_row, max_rows)
    cols = min(formula_ws.max_column, max_cols)
    lines = [
        f"### Sheet: {formula_ws.title} "
        f"(showing {rows}x{cols} of {formula_ws.max_row}x{formula_ws.max_column})"
    ]
    header = "\t".join([""] + [get_column_letter(c) for c in range(1, cols + 1)])
    lines.append(header)
    for r in range(1, rows + 1):
        cells = []
        for c in range(1, cols + 1):
            cells.append(format_cell(formula_ws.cell(row=r, column=c), value_ws.cell(row=r, column=c)))
        lines.append("\t".join([str(r)] + cells))
    if formula_ws.max_row > rows:
        lines.append(f"... {formula_ws.max_row - rows} more row(s) truncated")
    if formula_ws.max_column > cols:
        lines.append(f"... {formula_ws.max_column - cols} more column(s) truncated")
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", help="path to the workbook")
    parser.add_argument("--max-rows", type=int, default=DEFAULT_MAX_ROWS)
    parser.add_argument("--max-cols", type=int, default=DEFAULT_MAX_COLS)
    args = parser.parse_args()

    formula_wb = load_workbook(args.xlsx, data_only=False)
    value_wb = load_workbook(args.xlsx, data_only=True)

    parts = []
    for sheet_name in formula_wb.sheetnames:
        parts.append(
            dump_sheet(
                formula_wb[sheet_name],
                value_wb[sheet_name],
                args.max_rows,
                args.max_cols,
            )
        )
    print("\n\n".join(parts))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # noqa: BLE001 - surface any failure to the caller
        print(f"dump_workbook: {exc}", file=sys.stderr)
        sys.exit(1)
