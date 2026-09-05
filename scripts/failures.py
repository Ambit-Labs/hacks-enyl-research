#!/usr/bin/env python3
"""Turn evaluator output into a failure list grouped by bucket.

Usage:
    uv run --project eval python scripts/failures.py <run-dir> [--dataset-dir DIR]

Reads <run-dir>/results.json (written by scripts/score.sh), <run-dir>/predictions.jsonl,
and <run-dir>/traces/<id>.jsonl (optional; issue #6), and writes <run-dir>/failures.md
with counts by instruction type and by failure kind, followed by one row per failed task.

Never reads golden workbooks or golden values. Only predicted output workbooks and the
evaluator's own verdicts in results.json are used to classify a failure.
"""

import argparse
import json
import os
import sys
import time
from pathlib import Path

SCRIPT_DIR = Path(__file__).resolve().parent
REPO_ROOT = SCRIPT_DIR.parent
sys.path.insert(0, str(REPO_ROOT / "eval"))

import openpyxl  # noqa: E402

import sb  # noqa: E402

try:
    from rich.console import Console
    from rich.progress import BarColumn, Progress, TextColumn, TimeElapsedColumn, TimeRemainingColumn

    HAVE_RICH = True
except ImportError:
    HAVE_RICH = False

FAILURE_KINDS = ["no submit", "error value in cell", "wrong value", "timeout", "exception"]

EXCEL_ERROR_VALUES = {
    "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#SPILL!", "#CALC!",
}


def eprint(*args, **kwargs):
    print(*args, file=sys.stderr, **kwargs)


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("run_dir", help="a predict.ts output directory, e.g. runs/smoke")
    p.add_argument("--dataset-dir", default=os.environ.get("SB_DATASET_DIR"),
                    help="defaults to $SB_DATASET_DIR")
    return p.parse_args()


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8") as f:
        return [json.loads(line) for line in f if line.strip()]


def is_failed(item: dict) -> bool:
    if item["status"] == "no_golden":
        return False
    if item["status"] == "graded" and item.get("pass"):
        return False
    return True


def trace_info(run_dir: Path, task_id: str) -> tuple[str | None, str | None]:
    """Return (model_calls, submit_called) as strings, or (None, None) if traces aren't available."""
    traces_dir = run_dir / "traces"
    if not traces_dir.is_dir():
        return None, None
    trace_path = traces_dir / f"{task_id}.jsonl"
    if not trace_path.exists():
        return None, None
    calls = read_jsonl(trace_path)
    submitted = any(call.get("tool") == "submit" for call in calls)
    return str(len(calls)), ("yes" if submitted else "no")


def error_cells_in_output(task: dict, output_path: Path) -> list[str]:
    """Answer cells in the predicted output that hold an Excel error value."""
    if not output_path.exists():
        return []
    try:
        wb = openpyxl.load_workbook(output_path, data_only=True)
    except Exception:
        return []
    found = []
    for sheet, coord in sb.answer_cells(task, wb):
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        value = ws[coord].value
        if isinstance(value, str) and value.strip().upper() in EXCEL_ERROR_VALUES:
            found.append(f"{ws.title}!{coord}")
    return found


def classify(item: dict, prediction: dict | None, submit_called: str | None, error_cells: list[str]) -> str:
    """Pick one of FAILURE_KINDS for a failed task.

    Priority: a run-time signal (timeout, exception) beats an output-quality signal
    (no submit, error value, wrong value), because it explains why the output is bad.
    """
    pred_status = prediction.get("status") if prediction else None

    if pred_status == "timeout":
        return "timeout"
    if pred_status in ("model_failed", "error") or item["status"] == "error":
        return "exception"
    # missing_output at predict time means predict.ts never got a real output file from
    # the agent and fell back to copying the init workbook; that is a no-submit failure
    # even though the evaluator sees a gradeable file. A missing prediction line, or an
    # evaluator-side missing_output/missing, means the task was never produced at all.
    if pred_status == "missing_output" or prediction is None or item["status"] in ("missing", "missing_output"):
        return "no submit"
    if submit_called == "no":
        return "no submit"
    if error_cells:
        return "error value in cell"
    return "wrong value"


def make_progress(total: int):
    if HAVE_RICH and sys.stderr.isatty():
        progress = Progress(
            TextColumn("[progress.description]{task.description}"),
            BarColumn(),
            TextColumn("{task.completed}/{task.total}"),
            TimeElapsedColumn(),
            TimeRemainingColumn(),
            console=Console(stderr=True),
        )
        progress.start()
        task_ref = progress.add_task("failures.py", total=total)

        def advance(n: int, _outcomes: dict[str, int]):
            progress.update(task_ref, completed=n)

        def close():
            progress.stop()

        return advance, close

    start = time.monotonic()

    def advance(n: int, outcomes: dict[str, int]):
        if n % 20 == 0 or n == total:
            elapsed = time.monotonic() - start
            counts = " ".join(f"{k}={v}" for k, v in outcomes.items())
            eprint(f"failures.py: {n}/{total} tasks classified ({elapsed:.0f}s elapsed) {counts}")

    def close():
        pass

    return advance, close


def main() -> None:
    args = parse_args()
    run_dir = Path(args.run_dir)
    results_path = run_dir / "results.json"
    predictions_path = run_dir / "predictions.jsonl"

    if not results_path.exists():
        eprint(f"failures.py: no {results_path}; run scripts/score.sh {run_dir} first")
        sys.exit(1)
    if not args.dataset_dir:
        eprint("failures.py: --dataset-dir is required (or set SB_DATASET_DIR)")
        sys.exit(1)

    results = json.loads(results_path.read_text())
    predictions = {p["id"]: p for p in read_jsonl(predictions_path)}
    dataset_by_id = {t["id"]: t for t in sb.load_dataset(args.dataset_dir)}

    # results.json holds every dataset task when score.sh's --all flag scored the whole
    # dataset; a task this run never attempted isn't a bucket-worthy failure, so scope
    # the report to the ids actually present in predictions.jsonl.
    scoped_items = [item for item in results["items"] if item["id"] in predictions]
    failed = [item for item in scoped_items if is_failed(item)]
    eprint(f"failures.py: {run_dir}: {len(scoped_items)} scored task(s) in this run, {len(failed)} failed; classifying")

    rows = []
    outcomes = {kind: 0 for kind in FAILURE_KINDS}
    by_type = {}
    advance, close_progress = make_progress(len(failed))

    for i, item in enumerate(failed, start=1):
        task_id = item["id"]
        task = dataset_by_id.get(task_id)
        prediction = predictions.get(task_id)
        model_calls, submit_called = trace_info(run_dir, task_id)

        error_cells = []
        if task is not None and prediction is not None and prediction.get("output"):
            output_path = sb.resolve_output(prediction["output"], str(predictions_path))
            if output_path is not None:
                error_cells = error_cells_in_output(task, output_path)

        kind = classify(item, prediction, submit_called, error_cells)
        outcomes[kind] += 1
        instruction = task["instruction"][:120] if task else ""
        by_type[item["type"]] = by_type.get(item["type"], 0) + 1

        rows.append({
            "id": task_id,
            "type": item["type"],
            "status": item["status"],
            "model_calls": model_calls if model_calls is not None else "n/a",
            "submit_called": submit_called if submit_called is not None else "n/a",
            "error_cells": ", ".join(error_cells) if error_cells else "-",
            "kind": kind,
            "instruction": instruction,
        })
        advance(i, outcomes)

    close_progress()

    write_report(run_dir, len(scoped_items), failed, rows, outcomes, by_type)
    eprint(f"failures.py: wrote {run_dir / 'failures.md'}")


def write_report(run_dir: Path, total_scored: int, failed: list[dict], rows: list[dict],
                  outcomes: dict[str, int], by_type: dict[str, int]) -> None:
    traces_dir = run_dir / "traces"
    lines = [
        f"# Failures: {run_dir}",
        "",
        f"{len(failed)} of {total_scored} scored tasks failed.",
        "",
    ]
    if not traces_dir.is_dir():
        lines += [f"No `{traces_dir}` directory found; model call counts and submit status are n/a.", ""]

    lines += ["## By instruction type", "", "| type | failed |", "| --- | --- |"]
    for t, count in sorted(by_type.items()):
        lines.append(f"| {t} | {count} |")
    lines += [f"| **total** | **{sum(by_type.values())}** |", ""]

    lines += ["## By failure kind", "", "| kind | failed |", "| --- | --- |"]
    for kind in FAILURE_KINDS:
        lines.append(f"| {kind} | {outcomes[kind]} |")
    lines += [f"| **total** | **{sum(outcomes.values())}** |", ""]

    lines += [
        "## Failed tasks",
        "",
        "| id | type | status | kind | model calls | submit | error cells | instruction (first 120 chars) |",
        "| --- | --- | --- | --- | --- | --- | --- | --- |",
    ]
    for row in rows:
        instruction = row["instruction"].replace("|", "\\|").replace("\n", " ")
        lines.append(
            f"| {row['id']} | {row['type']} | {row['status']} | {row['kind']} | "
            f"{row['model_calls']} | {row['submit_called']} | {row['error_cells']} | {instruction} |"
        )
    lines.append("")

    (run_dir / "failures.md").write_text("\n".join(lines))


if __name__ == "__main__":
    main()
